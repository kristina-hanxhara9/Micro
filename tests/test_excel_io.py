from openpyxl import Workbook

from src.excel_io import OUTPUT_COLUMNS, read_retailer_names, write_results
from src.models import PriceSample, RetailerRecord


def _make_input(path, headers, rows):
    wb = Workbook()
    ws = wb.active
    if headers:
        ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_read_with_named_column(tmp_path):
    p = tmp_path / "in.xlsx"
    _make_input(p, ["name", "notes"], [["Best Buy", "x"], ["Target", "y"], [None, None]])
    assert read_retailer_names(p) == ["Best Buy", "Target"]


def test_read_no_header_uses_first_column(tmp_path):
    p = tmp_path / "in.xlsx"
    _make_input(p, None, [["Sephora"], ["B&H Photo"]])
    # First row will be treated as header; if header lookup fails it falls
    # back to col 0 starting from that row.
    names = read_retailer_names(p)
    assert "B&H Photo" in names


def test_write_splits_by_threshold(tmp_path):
    out = tmp_path / "out.xlsx"
    records = [
        RetailerRecord(
            name="High",
            website="https://high.com",
            confidence=0.85,
            categories=["a"],
            sample_prices=[PriceSample(item="x", price=1.0)],
        ),
        RetailerRecord(name="Low", confidence=0.3, flags=["no_website_found"]),
    ]
    n_ok, n_review = write_results(out, records, threshold=0.7)
    assert n_ok == 1
    assert n_review == 1

    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb["enriched"].max_row == 2  # header + 1
    assert wb["needs_review"].max_row == 2
    assert [c.value for c in wb["enriched"][1]] == OUTPUT_COLUMNS
