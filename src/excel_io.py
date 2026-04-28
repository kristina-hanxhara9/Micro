from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import RetailerRecord

OUTPUT_COLUMNS = [
    "name",
    "website",
    "about",
    "categories",
    "brands",
    "sample_prices",
    "phone",
    "email",
    "address",
    "confidence",
    "flags",
    "sources",
]


def read_retailer_names(path: str | Path, name_column: str = "name") -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []

    headers_lower = [str(h).strip().lower() if h is not None else "" for h in header]
    if name_column.lower() in headers_lower:
        idx = headers_lower.index(name_column.lower())
    else:
        # No matching header — treat first row as data, use column 0.
        idx = 0
        first_cell = header[0]
        names = [str(first_cell).strip()] if first_cell else []
        names.extend(str(r[idx]).strip() for r in rows if r and r[idx])
        return [n for n in names if n]

    return [str(r[idx]).strip() for r in rows if r and r[idx]]


def write_results(
    path: str | Path,
    records: list[RetailerRecord],
    threshold: float,
) -> tuple[int, int]:
    wb = Workbook()
    enriched = wb.active
    enriched.title = "enriched"
    review = wb.create_sheet("needs_review")

    enriched.append(OUTPUT_COLUMNS)
    review.append(OUTPUT_COLUMNS)

    n_enriched = n_review = 0
    for rec in records:
        row = _record_to_row(rec)
        if rec.confidence >= threshold:
            enriched.append(row)
            n_enriched += 1
        else:
            review.append(row)
            n_review += 1

    wb.save(path)
    return n_enriched, n_review


def _record_to_row(rec: RetailerRecord) -> list:
    return [
        rec.name,
        rec.website or "",
        rec.about or "",
        "; ".join(rec.categories),
        "; ".join(rec.brands),
        "; ".join(f"{p.item}: {p.price} {p.currency}" for p in rec.sample_prices),
        rec.phone or "",
        rec.email or "",
        rec.address or "",
        rec.confidence,
        "; ".join(rec.flags),
        "; ".join(rec.sources),
    ]
